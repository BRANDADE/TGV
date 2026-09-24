#!/usr/bin/env python3
import os
import sys
import csv
import argparse
import logging
import tempfile
import yaml

from scripts.core.vcf_loader import list_vcfs
from scripts.core.vcf_parser import parse_vcf_for_sample
from scripts.core.utils import get_analysis_prefix
from scripts.core.artifact_lookup import sample_id_from_vcf_name
from scripts.core.config_manager import get_safe_config_path
from scripts.core.analysis import build_run_trids, resolve_panel, run_analysis
from scripts.core.comments import LOW_COVERAGE, result_comments
from scripts.core.provenance import read_run_manifest, read_vcf_header, tgv_provenance

from scripts.bio.clinical_thresholds_loader import load_clinical_thresholds, load_trid_aliases
from scripts.bio.clinical_config_validator import validate_thresholds

# Minimal logging configuration for console output
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

# Optional import for Excel export
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# Codes retour (surveillés par cron / Slurm)
EXIT_OK = 0
EXIT_FAILURE = 1   # au moins un échantillon en échec, ou aucune ligne produite
EXIT_CONFIG = 2    # configuration ou entrée invalide : rien n'est écrit

# Ordre canonique de référence pour l'organisation des colonnes dans le fichier de sortie
CANONICAL_HEADERS = [
    "trgt_version",
    "bed_version",
    "run_id",
    "sample_id",
    "Locus",
    "Depth",
    "Genotype",
    "Classification",
    "Comments",
    # Provenance (E2) : de quoi reproduire chaque ligne
    "catalog",
    "thresholds_source",
    "karyotype",
    "tgv_version",
    "tgv_commit",
    "thresholds_sha256",
]

# Clé d'une ligne : relancer l'export d'un run remplace ses lignes au lieu de les dupliquer
ROW_KEY = ("run_id", "sample_id")


class ConfigError(Exception):
    """Configuration ou entrée invalide (code retour EXIT_CONFIG)."""


# ---------------------------------------------------------------------------
# Load Ataxia panel configuration from config_manager
# ---------------------------------------------------------------------------
def load_panel_ataxie():
    path = get_safe_config_path("buttons_panel.yaml")

    if not os.path.exists(path):
        raise ConfigError(f"Configuration file 'buttons_panel.yaml' not found at: {path}")

    try:
        with open(path, "r", encoding="utf-8") as f:
            data = yaml.safe_load(f)
    except Exception as e:
        raise ConfigError(f"Error reading buttons_panel.yaml: {e}")

    if not isinstance(data, dict) or "Ataxie" not in data:
        raise ConfigError("The panel section 'Ataxie' is missing from buttons_panel.yaml.")

    return data["Ataxie"]


# ---------------------------------------------------------------------------
# Header Alignment Utilities
# ---------------------------------------------------------------------------
def get_union_headers(existing_headers, new_headers):
    """
    Calcule l'union ordonnée des en-têtes existants et des nouveaux en-têtes
    en se basant sur l'ordre canonique défini.
    """
    union_set = set(existing_headers).union(set(new_headers))
    ordered = [h for h in CANONICAL_HEADERS if h in union_set]
    # Ajout d'éventuelles colonnes inattendues non définies dans CANONICAL_HEADERS
    for h in list(existing_headers) + list(new_headers):
        if h not in ordered:
            ordered.append(h)
    return ordered


def merge_rows(existing_rows, new_rows):
    """
    Lignes existantes privées de celles des (run_id, sample_id) réexportés, puis nouvelles lignes.
    """
    new_keys = {tuple(str(r.get(k, "")) for k in ROW_KEY) for r in new_rows}
    kept = [r for r in existing_rows if tuple(str(r.get(k, "") or "") for k in ROW_KEY) not in new_keys]
    removed = len(existing_rows) - len(kept)
    if removed:
        logging.info(f"Replacing {removed} existing row(s) of the same run/sample(s) in the output file.")
    return kept + list(new_rows)


def _atomic_target(out_path):
    """Fichier temporaire dans le même répertoire que la sortie (remplacement atomique)."""
    directory = os.path.dirname(os.path.abspath(out_path))
    fd, tmp_path = tempfile.mkstemp(prefix=".tgv_export_", dir=directory)
    os.close(fd)
    return tmp_path


# ---------------------------------------------------------------------------
# Flat Files Writers (CSV/TSV) with Header Re-alignment
# ---------------------------------------------------------------------------
def write_csv_tsv(out_path, rows, delimiter):
    existing_rows, existing_headers = [], []
    if os.path.exists(out_path) and os.path.getsize(out_path) > 0:
        with open(out_path, "r", newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            existing_headers = reader.fieldnames or []
            existing_rows = list(reader)

    headers = get_union_headers(existing_headers, list(rows[0].keys()))
    all_rows = merge_rows(existing_rows, rows)

    tmp_path = _atomic_target(out_path)
    try:
        with open(tmp_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers, delimiter=delimiter, extrasaction="ignore")
            writer.writeheader()
            for r in all_rows:
                writer.writerow({h: r.get(h, "") or "" for h in headers})
        os.replace(tmp_path, out_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ---------------------------------------------------------------------------
# Excel File Writer (XLSX) with Header Re-alignment
# ---------------------------------------------------------------------------
def write_xlsx(out_path, rows):
    if not HAS_OPENPYXL:
        raise ImportError(
            "The 'openpyxl' package is required to write Excel files (.xlsx). "
            "Please install it using 'pip install openpyxl' or choose '.csv' / '.tsv' output instead."
        )

    existing_rows, existing_headers = [], []
    if os.path.exists(out_path) and os.path.getsize(out_path) > 0:
        ws = openpyxl.load_workbook(out_path).active
        values = list(ws.iter_rows(values_only=True))
        if values:
            existing_headers = [h for h in values[0] if h is not None]
            for line in values[1:]:
                existing_rows.append({h: ("" if v is None else v) for h, v in zip(existing_headers, line)})

    headers = get_union_headers(existing_headers, list(rows[0].keys()))
    all_rows = merge_rows(existing_rows, rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRGT Export"
    ws.append(headers)
    for r in all_rows:
        ws.append([r.get(h, "") for h in headers])

    tmp_path = _atomic_target(out_path)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, out_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ---------------------------------------------------------------------------
# Write Rows to Output File (TSV, CSV, or XLSX)
# ---------------------------------------------------------------------------
def write_output_file(out_path, rows):
    if not rows:
        return

    _, ext = os.path.splitext(out_path.lower())

    if ext == ".tsv":
        write_csv_tsv(out_path, rows, delimiter="\t")
    elif ext == ".csv":
        write_csv_tsv(out_path, rows, delimiter=",")
    elif ext == ".xlsx":
        write_xlsx(out_path, rows)
    else:
        raise ValueError(
            f"Unsupported file extension '{ext}'. "
            "Supported formats are '.csv', '.tsv', and '.xlsx'."
        )


# ---------------------------------------------------------------------------
# One exported row per locus (raw values, no UI decoration)
# ---------------------------------------------------------------------------
def build_row(result, run_id, sample_id, low_depth_threshold, provenance):
    """
    provenance : trgt_version, bed_version (optionnel), catalog, karyotype,
                 tgv_version, tgv_commit, thresholds_sha256.
    """
    de = result.display_export
    comments = ["Low coverage" if c == LOW_COVERAGE else c for c in result_comments(result, low_depth_threshold)]

    row = {"trgt_version": provenance.get("trgt_version", "")}
    if provenance.get("bed_version") is not None:
        row["bed_version"] = provenance["bed_version"]
    row.update({
        "run_id": run_id,
        "sample_id": sample_id,
        "Locus": de.locus,
        "Depth": f"{de.depth1} / {de.depth2}",
        "Genotype": de.genotype,
        "Classification": de.classification,
        "Comments": "; ".join(comments),
        "catalog": provenance.get("catalog", ""),
        "thresholds_source": result.thresholds_source,
        "karyotype": provenance.get("karyotype", ""),
        "tgv_version": provenance.get("tgv_version", ""),
        "tgv_commit": provenance.get("tgv_commit", ""),
        "thresholds_sha256": provenance.get("thresholds_sha256", ""),
    })
    return row


def sample_provenance(zip_path, vcf_filename, sample_id, trgt_version_arg, bed_version, base, manifest):
    """
    Provenance d'un échantillon. La version de TRGT est lue dans l'en-tête du VCF
    (##trgtVersion) ; une valeur --trgt-version différente est une erreur.
    """
    header = read_vcf_header(zip_path, vcf_filename)
    trgt_version = header["trgt_version"]
    if trgt_version_arg and trgt_version and trgt_version_arg != trgt_version:
        raise ValueError(
            f"--trgt-version {trgt_version_arg} differs from the VCF header (##trgtVersion={trgt_version})"
        )

    provenance = dict(base)
    provenance.update({
        "trgt_version": trgt_version or (trgt_version_arg or ""),
        "bed_version": bed_version,
        "catalog": header["catalog"],
        "karyotype": ((manifest.get("samples") or {}).get(sample_id) or {}).get("karyotype", ""),
    })
    return provenance


# ---------------------------------------------------------------------------
# Export Pipeline for the Entire Run
# ---------------------------------------------------------------------------
def export_run(zip_path, trgt_version, bed_version, run_id, all_loci=False):
    """
    Analyse tous les patients du run.
    Retourne (rows, failures) ; failures = [(sample_id, message)].
    Lève ConfigError si la configuration ou l'archive est inutilisable.
    """
    logging.info(f"Opening TRGT archive: {zip_path}")

    if not os.path.isfile(zip_path):
        raise ConfigError(f"ZIP archive not found: {zip_path}")

    vcfs = list_vcfs(zip_path)
    if not vcfs:
        raise ConfigError("The ZIP archive does not contain any .trgt.vcf files.")

    # Load and validate clinical thresholds configuration
    thresholds_data = load_clinical_thresholds()
    errors, warnings = validate_thresholds(thresholds_data)
    for w in warnings:
        logging.warning(f"clinical_thresholds.yaml: {w}")
    if errors:
        raise ConfigError("Invalid clinical_thresholds.yaml: " + " | ".join(errors))

    label_priority = thresholds_data["label_priority"]
    low_depth_threshold = thresholds_data.get("low_depth_threshold", None)

    run_id = run_id or get_analysis_prefix(zip_path)
    trids, _, run_trids = build_run_trids(zip_path, thresholds_data, load_trid_aliases())

    base_provenance = tgv_provenance()
    manifest = read_run_manifest(zip_path)
    logging.info(
        f"TGV {base_provenance['tgv_version']} (commit {base_provenance['tgv_commit']}) | "
        f"clinical_thresholds.yaml SHA-256 {base_provenance['thresholds_sha256']}"
    )

    # Set up selected loci list depending on the choice of panel vs. all loci
    if all_loci:
        logging.info("Global export mode active: all detected loci will be processed.")
        requested = list(trids)
    else:
        panel_ataxie = load_panel_ataxie()
        requested, absent = resolve_panel(panel_ataxie, run_trids)
        logging.info(f"Ataxia clinical panel loaded ({len(panel_ataxie)} target loci, {len(requested)} in catalog).")
        if absent:
            logging.warning(f"Panel loci absent from the TRGT catalog: {absent}")

    rows = []
    failures = []

    # Sequential processing per sample (patient)
    for vcf_filename in vcfs:
        sample_id = sample_id_from_vcf_name(vcf_filename)
        logging.info(f"Analyzing patient: {sample_id}")

        try:
            provenance = sample_provenance(
                zip_path, vcf_filename, sample_id, trgt_version, bed_version, base_provenance, manifest,
            )
            sample_trids = parse_vcf_for_sample(zip_path=zip_path, vcf_filename=vcf_filename, global_trids=run_trids)
            results, discordances, missing = run_analysis(
                vcf_filename, sample_trids, run_trids, requested, label_priority, low_depth_threshold,
            )
        except Exception as e:
            logging.error(f"Analysis failed for sample {sample_id}: {e}", exc_info=True)
            failures.append((sample_id, str(e)))
            continue

        if discordances:
            logging.warning(f"Motif discordances (BED vs clinical YAML) for {sample_id}: {discordances}")
        if missing:
            logging.warning(f"Loci missing from the VCF of {sample_id}: {missing}")
        if not results:
            logging.warning(f"No targeted loci identified for sample {sample_id}")

        for result in results:
            rows.append(build_row(result, run_id, sample_id, low_depth_threshold, provenance))

    return rows, failures


# ---------------------------------------------------------------------------
# Main Execution Block
# ---------------------------------------------------------------------------
def main(argv=None):
    parser = argparse.ArgumentParser(description="Structured export for TRGT run analysis (command-line utility).")
    parser.add_argument("--zip", required=True, help="Path to the trgt_vcfs.zip archive")
    parser.add_argument("--trgt-version", default=None, help="Expected TRGT version (checked against ##trgtVersion of each VCF)")
    parser.add_argument("--bed-version", default=None, help="Clinical BED file version used")
    parser.add_argument("--run-id", default=None, help="Unique identifier for the run (default: ZIP name prefix)")
    parser.add_argument("--out", default="tgv_export_ataxie.tsv", help="Output filename (must end in .csv, .tsv, or .xlsx)")
    parser.add_argument("--all-loci", action="store_true", help="Export all detected loci instead of filtering for the clinical (Ataxia) panel only")
    args = parser.parse_args(argv)

    # Pre-validation of output format and dependencies
    _, ext = os.path.splitext(args.out.lower())
    if ext not in [".csv", ".tsv", ".xlsx"]:
        parser.error(
            f"Unsupported output file format '{ext}'. "
            "Please specify a filename ending in '.csv', '.tsv', or '.xlsx'."
        )

    if ext == ".xlsx" and not HAS_OPENPYXL:
        parser.error(
            "The 'openpyxl' package is required for Excel exports (.xlsx) but is not installed. "
            "Please install it (pip install openpyxl) or use '.csv' / '.tsv' output instead."
        )

    try:
        rows, failures = export_run(
            zip_path=args.zip,
            trgt_version=args.trgt_version,
            bed_version=args.bed_version,
            run_id=args.run_id,
            all_loci=args.all_loci,
        )
        if rows:
            write_output_file(args.out, rows)
            logging.info(f"{len(rows)} rows successfully exported to: {args.out}")
    except ConfigError as e:
        logging.error(f"Data export aborted: {e}")
        return EXIT_CONFIG
    except Exception as e:
        logging.error(f"Data export failed: {e}", exc_info=True)
        return EXIT_FAILURE

    if failures:
        for sample_id, message in failures:
            print(f"FAILED sample {sample_id}: {message}", file=sys.stderr)
        return EXIT_FAILURE

    if not rows:
        logging.error("No rows could be generated for export.")
        return EXIT_FAILURE

    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
